"""Regenera la sección «En LinkedIn y en medios» de la fuente de una página a partir de linkedin_actualizable.xlsx.

Uso: python build_medios.py linkedin_actualizable.xlsx fuente-una-pagina.html
Después: python build_pages.py fuente-una-pagina.html dist
"""
import sys, re, html
import openpyxl

xlsx, fuente = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(xlsx, data_only=True)

MES_ES = ["ene.", "feb.", "mar.", "abr.", "may.", "jun.", "jul.", "ago.", "sep.", "oct.", "nov.", "dic."]
MES_EN = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
ENTIDAD_EN = {"Escuela Politécnica de Cuenca": "the Polytechnic School of Cuenca",
              "UCLM": "UCLM", "Universidad de Castilla-La Mancha": "the University of Castilla-La Mancha",
              "SEEIC": "SEEIC"}
TIPO_EN = {"Prensa": "Press", "Radio": "Radio", "Televisión": "TV", "Entrevista": "Interview",
           "Premio": "Award", "Institucional": "Institutional"}
avisos = []

def e(t):
    return html.escape(str(t).strip(), quote=False)

def bi(es, en):
    return f'<span data-l="es">{es}</span><span data-l="en">{en}</span>'

def si(v):
    return str(v or "").strip().lower() in ("sí", "si", "s", "yes")

def fecha(v):
    """Devuelve (clave de orden, texto ES, texto EN, texto ES numérico)."""
    if hasattr(v, "year"):
        d, m, a = v.day, v.month, v.year
    else:
        p = [int(x) for x in re.findall(r"\d+", str(v))]
        if len(p) == 3:
            d, m, a = p
        elif len(p) == 2:
            d, (m, a) = 0, p
        else:
            raise ValueError(f"Fecha no válida: {v!r}")
    if d:
        return (a, m, d), f"{d} {MES_ES[m-1]} {a}", f"{d} {MES_EN[m-1]} {a}", f"{d:02d}/{m:02d}/{a}"
    # sin día: se ordena tras las del mismo mes con día
    return (a, m, 0), f"{MES_ES[m-1]} {a}", f"{MES_EN[m-1]} {a}", f"{m:02d}/{a}"

def en_o_es(en, es, que):
    if en and str(en).strip():
        return e(en)
    avisos.append(f"Falta la traducción al inglés de {que}: se usa el texto en español.")
    return e(es)

A = ' rel="noopener" target="_blank"'

# ---------- LinkedIn ----------
ws = wb["Publicaciones LinkedIn"]
cab = [c.value for c in ws[2]]
col = {k: i for i, k in enumerate(cab)}
def v(row, empieza):
    for k, i in col.items():
        if k and k.startswith(empieza):
            return row[i]
    return None

posts, destacada = [], None
for row in ws.iter_rows(min_row=3, values_only=True):
    if not v(row, "Fecha") or not v(row, "Enlace") or not si(v(row, "Mostrar")):
        continue
    orden, f_es, f_en, _ = fecha(v(row, "Fecha"))
    aut = (v(row, "Autoría") or "Propia").strip()
    ent = (v(row, "Entidad") or "").strip()
    if aut == "De otra entidad" and ent:
        tipo = bi(f"Publicación de la {e(ent)}" if ent.startswith("Escuela") or ent.startswith("Universidad") else f"Publicación de {e(ent)}",
                  f"Post by {e(ENTIDAD_EN.get(ent, ent))}")
    elif aut == "Compartida con comentario":
        tipo = bi("Compartida con comentario", "Shared with comment")
    elif aut == "Mención o etiqueta":
        tipo = bi("Mención", "Mention")
    else:
        tipo = bi("Publicación propia", "Own post")
    p = dict(orden=orden, f=bi(f_es, f_en), tipo=tipo, url=e(v(row, "Enlace")),
             tit_es=v(row, "Titular (solo"), tit_en=v(row, "Titular en inglés"),
             res_es=v(row, "Resumen en español"), res_en=v(row, "Resumen en inglés"))
    if si(v(row, "Destacada")):
        if destacada:
            avisos.append("Hay más de una publicación destacada: se usa la más reciente.")
            if p["orden"] < destacada["orden"]:
                posts.append(p); continue
            posts.append(destacada)
        destacada = p
    else:
        posts.append(p)

posts.sort(key=lambda p: p["orden"], reverse=True)

def resumen(p):
    return bi(e(p["res_es"]), en_o_es(p["res_en"], p["res_es"], f"la publicación de {p['orden']}"))

out = '<h3 id="linkedin">' + bi("Desde LinkedIn", "From LinkedIn") + '</h3>'
out += '<!-- Sección generada con build_medios.py a partir de linkedin_actualizable.xlsx: cambia el Excel, no este bloque. -->'
if destacada:
    d = destacada
    if not d["tit_es"]:
        avisos.append("La publicación destacada no tiene titular.")
    out += ('<article class="li-featured"><div><p class="kick"><b>' + bi("Destacada", "Featured") + '</b>' +
            bi("Publicación en LinkedIn", "LinkedIn post") + ' · ' + d["f"] + '</p>' +
            '<h4>' + bi(e(d["tit_es"] or ""), en_o_es(d["tit_en"], d["tit_es"] or "", "el titular destacado")) + '</h4>' +
            '<p>' + resumen(d) + '</p></div><div class="acts"><a class="btn li" href="' + d["url"] + '"' + A + '>' +
            bi("Leer en LinkedIn", "Read on LinkedIn") + '</a></div></article>')
out += '<div class="posts">'
for p in posts:
    tit = ''
    if p["tit_es"]:
        tit = '<p><strong>' + bi(e(p["tit_es"]), en_o_es(p["tit_en"], p["tit_es"], "un titular")) + '</strong></p>'
    out += ('<div class="post"><p class="meta">' + p["f"] + ' · ' + p["tipo"] + '</p>' + tit +
            '<p>' + resumen(p) + '</p><p><a href="' + p["url"] + '"' + A + '>' +
            bi("Ver publicación en LinkedIn", "View post on LinkedIn") + '</a></p></div>')
out += '</div>'
out += ('<p class="acts"><a class="btn li" href="https://www.linkedin.com/in/cesarsanchezmelendez" rel="me noopener" target="_blank">' +
        bi("Seguirme en LinkedIn", "Follow me on LinkedIn") + '</a></p>')
out += '<p class="nota">' + bi("Las publicaciones se enlazan sin incrustar para no cargar cookies de terceros.",
                               "Posts are linked, not embedded, to avoid loading third-party cookies.") + '</p>'

# ---------- Prensa ----------
ws = wb["Prensa y entrevistas"]
noticias = []
for row in ws.iter_rows(min_row=3, values_only=True):
    f, medio, prog, tipo, tit, url, mostrar = row[:7]
    if not f or not tit or not si(mostrar):
        continue
    orden, _, f_en, f_num = fecha(f)
    noticias.append((orden, f_num, f_en, medio, prog, tipo, tit, url))
noticias.sort(key=lambda n: n[0], reverse=True)

out += '<h3 id="prensa">' + bi("Entrevistas y prensa", "Interviews and press") + '</h3>'
out += '<p class="nota">' + bi("Los titulares se muestran en su idioma original (español). Los enlaces abren la fuente en una pestaña nueva.",
                               "Headlines are shown in their original language (Spanish). Links open the source in a new tab.") + '</p>'
out += '<ul class="news">'
for orden, f_num, f_en, medio, prog, tipo, tit, url in noticias:
    meta = bi(f_num, f_en) + ' · ' + e(medio) + (' · ' + e(prog) if prog else '') + ' · ' + bi(e(tipo), TIPO_EN.get(tipo, e(tipo)))
    enlace = (' · <a href="' + e(url) + '"' + A + '>' + bi("Leer o escuchar", "Read or listen") + '</a>') if url else '<span class="ph"> · [enlace pendiente]</span>'
    out += '<li><span class="meta">' + meta + '</span><span class="t">' + e(tit) + '</span>' + enlace + '</li>'
out += '</ul>'

# ---------- sustitución en la fuente ----------
S = open(fuente, encoding="utf-8").read()
i = S.index('<h3 id="linkedin">')
j = S.index('</section>', i)
S = S[:i] + out + S[j:]
open(fuente, "w", encoding="utf-8").write(S)
print(f"LinkedIn: {len(posts)} publicaciones + {'1 destacada' if destacada else 'sin destacada'}; prensa: {len(noticias)} entradas")
for a in avisos:
    print("AVISO:", a)
